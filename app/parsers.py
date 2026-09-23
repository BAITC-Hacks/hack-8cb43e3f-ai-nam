"""Read documents without executing content; keep stable source locations."""
import hashlib
import io
import re
import zipfile
from pathlib import Path

from .models import Clause, Document

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNPACKED_BYTES = 100 * 1024 * 1024
MAX_CLAUSES = 5000
FORMATS = {".docx", ".pdf", ".xlsx", ".txt"}
NUMBER = re.compile(r"^(\d{1,2}(?:\.\d{1,3}){0,5})[.)]?\s*(.*)$")
BULLET = re.compile(r"^([а-яa-z])[.)]\s+(.*)$", re.I)
ACTION = re.compile(r"организ|осуществ|обеспеч|готов|подгот|провод|провед|соглас|разраб|анализ|оценк|оцени|провер|контрол|выяв|определ|представ|вынос|запраш|запрос|утверж|формир|монитор|участв|взаимод|выполн|информир|вести |довод|консульт|распредел|принимать|использ|соблюд|руковод|отвеч|ұйымдастыр|қамтамасыз|дайындай|жүргіз|бекіт|келіс|бағалай|бақыла|талдай|тексер", re.I)
ROLE = re.compile(r"^(?:Главн\w+ аудитор|Директор\w*|Работник\w*|Руководител\w*|Начальник\w*|Отдел\w*|Департамент\w*|Служб\w*|Комитет\w*|Бас аудитор|Қызметкер\w*|Бөлім\w*|Басқарма\w*)", re.I)
UNIT = re.compile(r"^(?:Департамент|Отдел|Управление|Служба|Центр|Бюро|Бөлім|Басқарма|Қызмет|Орталық)\s+", re.I)


def clean(text):
    return re.sub(r"\s+", " ", text.replace("\xa0", " ").replace("\u00ad", "")).strip()


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED_BYTES:
            raise ValueError("Распакованный документ слишком большой. Разделите его на несколько файлов.")
        if any(info.flag_bits & 1 for info in archive.infolist()):
            raise ValueError("Документ защищён паролем. Загрузите незашифрованную копию.")


class WordNumbering:
    """Recover simple automatic Word list labels, including nested lists."""
    def __init__(self, doc):
        from docx.oxml.ns import qn
        self.qn = qn
        self.defs, self.counters = {}, {}
        try:
            root = doc.part.numbering_part.element
        except (KeyError, NotImplementedError):
            return
        abstracts = {}
        for abstract in root.findall(qn("w:abstractNum")):
            levels = {}
            for level in abstract.findall(qn("w:lvl")):
                def value(name, default):
                    el = level.find(qn("w:" + name))
                    return el.get(qn("w:val")) if el is not None else default
                levels[int(level.get(qn("w:ilvl")))] = (value("lvlText", "%1."), value("numFmt", "decimal"), int(value("start", "1")))
            abstracts[abstract.get(qn("w:abstractNumId"))] = levels
        for num in root.findall(qn("w:num")):
            abstract = num.find(qn("w:abstractNumId"))
            if abstract is not None:
                self.defs[num.get(qn("w:numId"))] = abstracts.get(abstract.get(qn("w:val")), {})

    def label(self, paragraph):
        props = paragraph._p.pPr
        if props is None or props.numPr is None or props.numPr.numId is None:
            return ""
        num_id = str(props.numPr.numId.val)
        level = props.numPr.ilvl.val if props.numPr.ilvl is not None else 0
        template, fmt, start = self.defs.get(num_id, {}).get(level, ("%1.", "decimal", 1))
        counters = self.counters.setdefault(num_id, {})
        counters[level] = counters.get(level, start - 1) + 1
        for deeper in list(counters):
            if deeper > level:
                del counters[deeper]
        if fmt == "bullet":
            return "•"
        for depth in range(9):
            number = counters.get(depth, 1)
            label = str(number)
            if depth == level and fmt in {"lowerLetter", "upperLetter"}:
                label = chr((97 if fmt == "lowerLetter" else 65) + (number - 1) % 26)
            elif depth == level and fmt == "russianLower":
                label = "абвгдежзийклмнопрстуфхцчшщэюя"[(number - 1) % 29]
            template = template.replace("%" + str(depth + 1), label)
        return template


def read_docx(data):
    from docx import Document as WordDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = WordDocument(io.BytesIO(data))
    numbering = WordNumbering(doc)
    blocks, warnings = [], []
    paragraph_count, table_count = 0, 0
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_count += 1
            text = block.text
            label = numbering.label(block)
            if label and not NUMBER.match(text) and not BULLET.match(text):
                text = label + " " + text
            if clean(text):
                blocks.append((text, f"Абзац {paragraph_count}"))
        elif isinstance(block, Table):
            table_count += 1
            for row_i, row in enumerate(block.rows, 1):
                seen = set()
                for col_i, cell in enumerate(row.cells, 1):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    if clean(cell.text):
                        blocks.append((cell.text, f"Таблица {table_count}, строка {row_i}, ячейка {col_i}"))
    if doc.inline_shapes or b"<w:drawing" in doc.part.blob:
        warnings.append("В Word есть изображения или схемы. Текст внутри них не распознан; проверьте оригинал.")
    return blocks, warnings


def read_pdf(data):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("PDF защищён паролем. Загрузите незашифрованную копию.")
    if len(reader.pages) > 300:
        raise ValueError("В PDF больше 300 страниц. Разделите его на части.")
    blocks, empty = [], []
    for page_i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if len(clean(text)) < 20:
            empty.append(str(page_i))
        # Preserve numbered paragraphs while joining line wraps.
        groups = re.split(r"\n(?=\s*(?:\d+(?:\.\d+)*[.)]|[а-яa-z][.)])\s*)|\n\s*\n", text)
        blocks.extend((x, f"Страница {page_i}") for x in groups if clean(x))
    warnings = []
    if empty:
        warnings.append("Страницы без читаемого текста: " + ", ".join(empty[:20]) + ". Для сканов нужен PDF с текстовым слоем (OCR).")
    return blocks, warnings


def read_xlsx(data):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    blocks = []
    try:
        for sheet in book:
            for row in sheet.iter_rows():
                values = [f"{cell.value}" for cell in row if cell.value is not None]
                if values:
                    filled = [cell for cell in row if cell.value is not None]
                    blocks.append((" | ".join(values), f"Лист «{sheet.title}», {filled[0].coordinate}:{filled[-1].coordinate}"))
                if len(blocks) > MAX_CLAUSES:
                    raise ValueError("В Excel слишком много строк. Оставьте листы с положениями и функциями.")
    finally:
        book.close()
    return blocks, ["Excel: анализируются значения ячеек. Схемы, изображения и формулы без сохранённого результата требуют проверки."]


def modality(text, inherited=""):
    t = (text + " " + inherited).lower()
    if re.search(r"не имеют права|не имеет права|запрещ|не вправе|тыйым салынады", t):
        return "forbidden"
    if re.search(r"может|могут|вправе|име[ею]т право|имеют право|құқылы|мүмкін", t):
        return "permission"
    if re.search(r"долж[еон]|обязан|необходимо|міндетті|тиіс", t):
        return "required"
    return "statement"


def segment(doc, blocks):
    parents, roles = {}, {}
    section, last_number, active_actor, actor_source = "", "", "", ""
    in_toc = False
    for raw, location in blocks:
        # Imported Word files sometimes contain several numbered clauses in one paragraph.
        parts = re.split(r"\s+(?=\d{1,2}\.\d+(?:\.\d+)*\.\s*[А-ЯA-Z])", clean(raw))
        for text in parts:
            if not text:
                continue
            if text.lower().strip(" .") == "оглавление":
                in_toc = True
            if in_toc:
                continue
            number, body, parent = "", text, None
            match = NUMBER.match(text)
            bullet = BULLET.match(text)
            if match:
                number, body = match.groups()
                last_number = number
                if "." not in number and len(body) < 160:
                    section = body
                    active_actor, actor_source = "", ""
                parent_num = number.rsplit(".", 1)[0] if "." in number else ""
                parent = parents.get(parent_num)
                ancestor = parent_num
                while ancestor:
                    if ancestor in roles:
                        active_actor, actor_source = roles[ancestor]
                        break
                    ancestor = ancestor.rsplit(".", 1)[0] if "." in ancestor else ""
            elif bullet:
                letter, body = bullet.groups()
                number = last_number + "." + letter.lower() if last_number else letter.lower()
                parent = parents.get(last_number)
            clause_id = f"{doc.id}-c{len(doc.clauses) + 1}"
            context = parent.text if parent else ""
            role_header = bool(ROLE.match(body) and body.rstrip().endswith(":") and (len(body) < 180 or "имеют право" in body or "имеет право" in body))
            # Sentences declaring a role's rights inherit that role, even when long.
            if ROLE.match(body) and re.search(r"(?:а также )?имеют? право:\s*$", body):
                role_header = True
            kind = "other"
            if role_header:
                actor = re.split(r" обязан| обязаны| имеют право| имеет право| не имеют права| не имеет права", body, flags=re.I)[0].rstrip(" :")
                active_actor, actor_source = actor, clause_id
                if number:
                    roles[number] = (actor, clause_id)
                kind = "role"
            elif UNIT.match(body) and ("подразделен" in context or "состоит" in context or (len(body) < 160 and not ACTION.search(body))):
                kind = "unit"
            elif bullet and ("подчиняются" in context or "должностей" in context):
                kind = "position"
            elif number and "." not in number:
                kind = "heading"
            elif (number or active_actor) and (ACTION.search(body) or (parent and parent.kind in {"function", "role"} and bullet)):
                kind = "function"
            if "термины и определения" in section.lower():
                kind = "other"
            if body.strip(" ;:.") == "":
                kind = "empty"
                doc.warnings.append(f"Пункт {number}: пустое содержание в исходном документе.")
            clause = Clause(clause_id, doc.id, number, text, location, section, active_actor, actor_source, context, parent.id if parent else "", kind, modality(body, context))
            doc.clauses.append(clause)
            if number and not bullet:
                parents[number] = clause
            if len(doc.clauses) > MAX_CLAUSES:
                raise ValueError("В документе больше 5000 фрагментов. Разделите его на части.")
    if not any(c.kind == "function" for c in doc.clauses):
        doc.warnings.append("Явные формулировки функций не выделены. Проверьте структуру и текст документа.")


def parse_document(data: bytes, name: str, side: str, doc_id: str) -> Document:
    suffix = Path(name).suffix.lower()
    if suffix not in FORMATS:
        raise ValueError("Поддерживаются DOCX, PDF, XLSX и TXT. Старые DOC/XLS сохраните в новом формате.")
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("Файл пуст или превышает 20 МБ.")
    if suffix in {".docx", ".xlsx"}:
        check_archive(data)
    try:
        if suffix == ".docx":
            blocks, warnings = read_docx(data)
        elif suffix == ".pdf":
            blocks, warnings = read_pdf(data)
        elif suffix == ".xlsx":
            blocks, warnings = read_xlsx(data)
        else:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1251")
            blocks = [(line, f"Строка {i}") for i, line in enumerate(text.splitlines(), 1) if clean(line)]
            warnings = []
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Не удалось прочитать {name}. Проверьте, что файл открывается и не защищён паролем.") from exc
    if not blocks:
        raise ValueError(f"В файле {name} нет читаемого текста. Для скана сначала выполните OCR.")
    doc = Document(doc_id, name, side, suffix[1:].upper(), len(data), hashlib.sha256(data).hexdigest(), warnings=warnings)
    segment(doc, blocks)
    return doc
