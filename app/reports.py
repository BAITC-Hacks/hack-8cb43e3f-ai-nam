import html
import io
from datetime import datetime

REVIEW = {"pending": "Не проверено", "confirmed": "Подтверждено", "rejected": "Отклонено"}


def docx_report(case):
    """Editable report. Every conclusion carries the original, unabridged evidence."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

    report = Document()
    section = report.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = report.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [('Title', 24), ('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11)]:
        style = report.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = 'Arial', Pt(size), RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(8)
    quote = report.styles['Quote']
    quote.font.name, quote.font.size, quote.font.italic = 'Arial', Pt(10.5), False
    quote.font.color.rgb = RGBColor.from_string('334155')
    quote.paragraph_format.left_indent = Inches(.15)
    quote.paragraph_format.right_indent = 0
    quote.paragraph_format.space_after = Pt(8)
    header = section.header.paragraphs[0]
    header.text = 'AI-NAM   /   Анализ организационных изменений'
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string('64748B')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Страница ').font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    report.core_properties.title = 'Заключение об организационных изменениях'
    report.core_properties.author = 'AI-NAM'
    report.core_properties.subject = case['title']
    result, stats = case['result'], case['result']['stats']
    docs = {d['id']: (i + 1, d) for i, d in enumerate(result['documents'])}
    sources = {c['id']: (d, c) for d in result['documents'] for c in d['clauses']}

    def ref(cid):
        d, c = sources[cid]
        return f'Д{docs[d["id"]][0]}, п. {c["number"] or "без номера"}, {c["location"]}'

    def paragraph(label, value):
        p = report.add_paragraph()
        p.add_run(label + ' ').bold = True
        p.add_run(str(value))
        return p

    def table(heads, rows, widths):
        tab = report.add_table(rows=1, cols=len(heads))
        tab.alignment, tab.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        for col, width in zip(tab.columns, widths):
            col.width = Inches(width)
        for row in [heads, *rows]:
            cells = tab.rows[0].cells if row is heads else tab.add_row().cells
            for cell, value, width in zip(cells, row, widths):
                cell.width, cell.text = Inches(width), str(value)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement('w:tcMar')
                for side in ['top', 'left', 'bottom', 'right']:
                    m = OxmlElement('w:' + side)
                    m.set(qn('w:w'), '100')
                    m.set(qn('w:type'), 'dxa')
                    margins.append(m)
                tcpr.append(margins)
                borders = OxmlElement('w:tcBorders')
                for side in ['top', 'left', 'bottom', 'right']:
                    b = OxmlElement('w:' + side)
                    for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                        b.set(qn('w:' + key), val)
                    borders.append(b)
                tcpr.append(borders)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(10)
                        if row is heads:
                            run.bold = True
                if row is heads:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'EEF1F5')
                    tcpr.append(shade)
        repeat = OxmlElement('w:tblHeader')
        tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        report.add_paragraph().paragraph_format.space_after = Pt(2)
        return tab

    report.add_paragraph('Заключение об организационных изменениях', 'Title')
    report.add_paragraph(case['title'])
    paragraph('Сформировано', datetime.now().strftime('%d.%m.%Y %H:%M'))
    paragraph('Автор сравнения', case.get('owner_name') or 'Не указан')
    checked = sum(f['review'] == 'confirmed' for f in result['findings'])
    pending = sum(f['review'] == 'pending' for f in result['findings'])
    report.add_paragraph(f'Сопоставлено {stats["matched"]} из {stats["before_functions"]} исходных фрагментов функций. Выявлено {stats["structure_changes"]} изменений в явных перечнях подразделений. Вопросов, требующих внимания: {stats["attention"]}. Аналитик подтвердил {checked} выводов; ожидают проверки {pending}.')
    report.add_paragraph('Заключение относится к загруженному комплекту «до» и «после». Автоматические выводы являются предварительными. Отсутствие соответствия в комплекте не доказывает утрату функции во всей организации. Перед утверждением проверьте источники и зафиксируйте решение по каждому спорному выводу.')
    paragraph('Метод', 'Структурный анализ и семантическая проверка ИИ' if result['method'] == 'hybrid' else 'Структурный анализ')
    if result.get('ai', {}).get('reviewed'):
        ai = result['ai']
        paragraph('Покрытие ИИ', f'{ai["reviewed"]} из {ai.get("eligible", 0)} спорных соответствий; модель {ai.get("model", "")}.')
    report.add_heading('Исходные документы', level=1)
    for number, d in docs.values():
        paragraph(f'Д{number}', f'{d["name"]} — комплект «{"до" if d["side"] == "before" else "после"}».')
        p = report.add_paragraph('SHA-256: ' + d['digest'])
        for run in p.runs:
            run.font.size = Pt(8)
        if d.get('source_url'):
            report.add_paragraph(d['source_url'])
    report.add_heading('Изменения подразделений', level=1)
    status = {'retained': 'Сохранено', 'added': 'Добавлено', 'missing': 'Не найдено'}
    if result['structure']:
        table(['Подразделение', 'Состояние', 'Источники'], [[s['name'], status[s['status']], '\n'.join(ref(cid) for cid in s['before_ids'] + s['after_ids'])] for s in result['structure']], [3.25, 1.05, 2.7])
    else:
        report.add_paragraph('Явный перечень подразделений не найден.')
    report.add_heading('Выводы и рекомендации', level=1)
    severity = {'high': 'Высокий', 'medium': 'Средний', 'low': 'Низкий', 'info': 'Информация'}
    for i, f in enumerate(sorted(result['findings'], key=lambda x: {'high': 0, 'medium': 1, 'low': 2, 'info': 3}[x['severity']]), 1):
        report.add_heading(f'{i} {f["label"]}', level=2)
        paragraph('Предмет', f['title'])
        paragraph('Приоритет и проверка', severity[f['severity']] + ' · ' + REVIEW[f['review']])
        report.add_paragraph(f['description'])
        for title, ids in [('До', f['before_ids']), ('После', f['after_ids'])]:
            p = paragraph(title, 'Явное соответствие не найдено.' if not ids else '')
            p.paragraph_format.keep_with_next = bool(ids)
            for cid in ids:
                p = report.add_paragraph(ref(cid))
                p.paragraph_format.keep_with_next = True
                for run in p.runs:
                    run.bold, run.font.size = True, Pt(9)
                report.add_paragraph(sources[cid][1]['text'], 'Quote')
        paragraph('Рекомендация', f['recommendation'])
        if f.get('ai_review'):
            paragraph('Комментарий ИИ', f['ai_review']['reason'])
        if f.get('comment'):
            paragraph('Комментарий аналитика', f['comment'])
        if f.get('reviewer'):
            paragraph('Проверил', f['reviewer'])
    if not result['findings']:
        report.add_paragraph('Значимых изменений не обнаружено.')
    report.add_heading('Матрица сопоставления функций', level=1)
    report.add_paragraph('Обозначения Д1, Д2 и далее соответствуют реестру документов выше. Полные тексты пар также доступны в таблице Excel и в приложении.')
    table(['Статус', 'До', 'После'], [[row['label'], '\n'.join(ref(cid) for cid in row['before_ids']) or '—', '\n'.join(ref(cid) for cid in row['after_ids']) or '—'] for row in result['matrix']], [1.45, 2.775, 2.775])
    report.add_heading('Ограничения и качество данных', level=1)
    for text in result['limitations']:
        report.add_paragraph(text, 'List Bullet')
    for d in result['documents']:
        for warning in d['warnings']:
            paragraph(d['name'], warning)
    output = io.BytesIO()
    report.save(output)
    return output.getvalue()


def html_report(case):
    esc = lambda x: html.escape(str(x or ""), quote=True)
    result = case["result"]
    sources = {c["id"]: {**c, "document": d["name"], "side": d["side"]} for d in result["documents"] for c in d["clauses"]}
    def refs(ids):
        return "".join(f'<p><a href="#{esc(cid)}">{esc(sources[cid]["document"])} · п. {esc(sources[cid]["number"] or "—")} · {esc(sources[cid]["location"])}</a></p><blockquote>{esc(sources[cid]["text"])}</blockquote>' for cid in ids)
    cards = []
    for f in sorted(result["findings"], key=lambda x: {"high": 0, "medium": 1, "low": 2, "info": 3}[x["severity"]]):
        cards.append(f'<article><div class="tag">{esc(f["label"])} · {esc(REVIEW[f["review"]])}</div><h3>{esc(f["title"])}</h3><p>{esc(f["description"])}</p><div class="columns"><section><h4>До</h4>{refs(f["before_ids"]) or "Нет исходного соответствия"}</section><section><h4>После</h4>{refs(f["after_ids"]) or "Явное соответствие не найдено"}</section></div><p><strong>Рекомендация:</strong> {esc(f["recommendation"])}</p>' + (f'<p><strong>Комментарий аналитика:</strong> {esc(f["comment"])}</p>' if f.get("comment") else "") + (f'<p>Проверил: {esc(f.get("reviewer"))}</p>' if f.get("reviewer") else "") + '</article>')
    structure = "".join(f'<tr><td>{esc(s["name"])}</td><td>{ {"retained":"Сохранено","added":"Добавлено","missing":"Не найдено"}[s["status"]] }</td></tr>' for s in result["structure"])
    docs = "".join(f'<h3>{esc(d["name"])}</h3><p>{esc("До" if d["side"] == "before" else "После")} · SHA-256: {esc(d["digest"])}</p>' + "".join(f'<p id="{esc(c["id"])}"><small>{esc(c["location"])}</small><br>{esc(c["text"])}</p>' for c in d["clauses"]) for d in result["documents"])
    stats = result["stats"]
    return f'''<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(case["title"])} — AI-NAM</title>
    <style>body{{font:15px/1.65 system-ui,sans-serif;color:#18212f;max-width:1100px;margin:40px auto;padding:0 25px}}h1{{font-size:32px}}h2{{margin-top:48px}}small,.tag{{color:#596879}}article{{border:1px solid #dce1e8;border-radius:12px;padding:24px;margin:20px 0;break-inside:avoid}}blockquote{{margin:0;border-left:3px solid #6961df;padding:12px;background:#f5f6fa;white-space:pre-wrap}}a{{color:#5145b9}}.columns{{display:grid;grid-template-columns:1fr 1fr;gap:20px}}table{{border-collapse:collapse;width:100%}}td,th{{border-bottom:1px solid #dde2e9;padding:10px;text-align:left}}.notice{{background:#f1f3fc;padding:20px;border-radius:12px}}@media print{{body{{margin:0;font-size:10pt}}h2{{break-before:page}}.no-print{{display:none}}a{{color:inherit;text-decoration:none}}}}@media(max-width:700px){{.columns{{grid-template-columns:1fr}}}}</style>
    <body><div class="tag">AI-NAM · АНАЛИТИЧЕСКОЕ ЗАКЛЮЧЕНИЕ</div><h1>{esc(case["title"])}</h1><p>Сформировано {datetime.now().strftime("%d.%m.%Y %H:%M")} · Автор сравнения: {esc(case.get("owner_name"))}</p>
    <p class="notice">Проанализировано документов: {stats["source_count"]}. Сопоставлено фрагментов функций: {stats["matched"]} из {stats["before_functions"]}. Вопросов для проверки: {stats["attention"]}. Метод: {"структурный анализ и ИИ" if result["method"] == "hybrid" else "структурный анализ"}. Выводы рекомендательные и относятся только к представленному комплекту.</p>
    <p class="no-print">Чтобы сохранить PDF: откройте меню печати браузера (Ctrl+P) и выберите «Сохранить как PDF».</p>
    <h2>Изменения подразделений</h2><table><tr><th>Подразделение</th><th>Состояние</th></tr>{structure or '<tr><td colspan="2">Явный перечень подразделений не найден.</td></tr>'}</table>
    <h2>Выводы и рекомендации</h2>{''.join(cards) or '<p>Значимых изменений не обнаружено.</p>'}
    <h2>Ограничения и качество данных</h2><ul>{''.join('<li>'+esc(x)+'</li>' for x in result['limitations'])}</ul>{''.join('<p>'+esc(d['name'])+': '+esc(w)+'</p>' for d in result['documents'] for w in d['warnings'])}
    <h2>Реестр исходных фрагментов</h2>{docs}</body></html>'''


def xlsx_report(case):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = Workbook()
    findings = book.active
    findings.title = "Выводы"
    matrix = book.create_sheet("Матрица функций")
    sources = book.create_sheet("Источники")
    structure = book.create_sheet("Подразделения")
    source_map = {c["id"]: (d, c) for d in case["result"]["documents"] for c in d["clauses"]}
    def safe(value):
        value = str(value or "")
        return "'" + value if value.startswith(("=", "+", "-", "@")) else value
    def refs(ids):
        return "\n".join(f'{source_map[cid][0]["name"]}, п. {source_map[cid][1]["number"]}, {source_map[cid][1]["location"]}: {source_map[cid][1]["text"]}' for cid in ids)
    findings.append(["Тип", "Вывод", "Обоснование", "До: источники", "После: источники", "Рекомендация", "Проверка", "Комментарий", "Проверил"])
    for f in case["result"]["findings"]:
        findings.append([safe(v) for v in [f["label"], f["title"], f["description"], refs(f["before_ids"]), refs(f["after_ids"]), f["recommendation"], REVIEW[f["review"]], f.get("comment"), f.get("reviewer")]])
    matrix.append(["Статус", "Функция до", "Функция после", "Сходство формулировок"])
    for row in case["result"]["matrix"]:
        matrix.append([safe(row["label"]), safe(refs(row["before_ids"])), safe(refs(row["after_ids"])), row["score"]])
    sources.append(["ID", "Комплект", "Документ", "Пункт", "Расположение", "Исполнитель", "Текст"])
    for cid, (doc, c) in source_map.items():
        sources.append([safe(v) for v in [cid, "До" if doc["side"] == "before" else "После", doc["name"], c["number"], c["location"], c["actor"], c["text"]]])
    structure.append(["Подразделение", "Статус", "Источники до", "Источники после"])
    for item in case["result"]["structure"]:
        structure.append([safe(v) for v in [item["name"], item["status"], refs(item["before_ids"]), refs(item["after_ids"])]])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="4F46A5")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 32
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 40 if column[0].column > 1 else 24
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
