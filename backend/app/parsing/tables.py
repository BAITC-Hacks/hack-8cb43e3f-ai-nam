"""Табличные источники (Excel, CSV, таблицы Word): штатные расписания, матрицы функций, оргструктуры.

Заголовок таблицы распознаётся по ключевым словам; строки превращаются в
пункты-записи со ссылкой «Лист, строка N» и полями unit/parent/function/position/count.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from .segment import Clause

COLUMN_RULES: dict[str, re.Pattern[str]] = {
    "code": re.compile(r"^(№|n|no|код|п/п|№ п/п|шифр)\b|^№", re.I),
    "parent": re.compile(r"вышестоящ|родител|подчин|входит в|в составе|parent|жоғары", re.I),
    "unit": re.compile(r"подразделени|структурн\w* единиц|департамент|управлени|отдел|бөлімше|unit", re.I),
    "function": re.compile(r"функци|задач|обязанност|полномочи|функция|міндет", re.I),
    "position": re.compile(r"должност|лауазым|position", re.I),
    "count": re.compile(r"численн|кол-?во|количеств|штатн\w* ед|единиц|саны", re.I),
    "head": re.compile(r"руководител|басшы", re.I),
}


def _detect_header(rows: list[list[str]]) -> tuple[int, dict[str, int]] | None:
    best: tuple[int, dict[str, int]] | None = None
    for r_idx, row in enumerate(rows[:12]):
        mapping: dict[str, int] = {}
        for c_idx, cell in enumerate(row):
            cell = (cell or "").strip()
            if not cell or len(cell) > 80:
                continue
            for key, rx in COLUMN_RULES.items():
                if key not in mapping and rx.search(cell):
                    # «Вышестоящее подразделение» — это parent, а не unit
                    if key == "unit" and COLUMN_RULES["parent"].search(cell):
                        continue
                    mapping[key] = c_idx
                    break
        if ("unit" in mapping or "function" in mapping) and len(mapping) >= 2:
            if best is None or len(mapping) > len(best[1]):
                best = (r_idx, mapping)
    return best


def rows_to_clauses(rows: list[list[str]], sheet: str, start_id: int = 1) -> tuple[list[Clause], list[dict[str, Any]]]:
    """Возвращает (пункты, записи). Записи используются при извлечении оргструктуры."""
    clauses: list[Clause] = []
    records: list[dict[str, Any]] = []
    header = _detect_header(rows)
    if header is None:
        for r_idx, row in enumerate(rows, start=1):
            cells = [c.strip() for c in row if c and c.strip()]
            if not cells:
                continue
            ref = f"{sheet}, стр. {r_idx}"
            clauses.append(Clause(id=f"c{start_id + len(clauses)}", ref=ref, ref_display=ref,
                                  text=" | ".join(cells), level=1, kind="row", source=ref))
        return clauses, records

    h_idx, cols = header
    last_unit = ""
    for r_idx in range(h_idx + 1, len(rows)):
        row = rows[r_idx]

        def get(key: str) -> str:
            i = cols.get(key)
            if i is None or i >= len(row):
                return ""
            return (row[i] or "").strip()

        unit = get("unit") or ""
        function = get("function")
        position = get("position")
        if not unit and (function or position):
            unit = last_unit  # объединённые ячейки: подразделение указано один раз
        if not any((unit, function, position)):
            continue
        last_unit = unit or last_unit
        parent = get("parent")
        count_raw = get("count")
        try:
            count = float(count_raw.replace(",", ".")) if count_raw else None
        except ValueError:
            count = None
        ref = f"{sheet}, стр. {r_idx + 1}"
        if function:
            text, kind = function, "function"
        elif position:
            text, kind = f"{position}" + (f" — {count_raw} ед." if count_raw else ""), "structure"
        else:
            text = unit + (f" (подчинено: {parent})" if parent else "")
            kind = "structure"
        rec = {
            "unit": unit, "parent": parent, "function": function, "position": position,
            "count": count, "head": get("head"), "code": get("code"), "ref": ref,
        }
        c = Clause(id=f"c{start_id + len(clauses)}", ref=ref, ref_display=ref, text=text, level=1,
                   kind=kind, actor=unit, section=sheet, source=ref, extra=rec)
        rec["clause_id"] = c.id
        clauses.append(c)
        records.append(rec)
    return clauses, records


def read_xlsx(path: Path) -> tuple[list[Clause], list[dict[str, Any]], dict]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    clauses: list[Clause] = []
    records: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c.strip() for c in r)] if rows else []
        cl, rec = rows_to_clauses(rows, ws.title, start_id=len(clauses) + 1)
        clauses.extend(cl)
        records.extend(rec)
    wb.close()
    return clauses, records, {"sheets": len(wb.sheetnames)}


def read_csv(path: Path) -> tuple[list[Clause], list[dict[str, Any]], dict]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "utf-8"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
    rows = list(csv.reader(text.splitlines(), dialect))
    clauses, records = rows_to_clauses(rows, path.stem)
    return clauses, records, {}
