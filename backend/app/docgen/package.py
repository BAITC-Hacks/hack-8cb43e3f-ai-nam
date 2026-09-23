"""Автоматическая подготовка пакета документов по оргструктуре.

Для каждого подразделения формируется проект «Положения о структурном подразделении»
(RU/KZ) с функциями из исходных документов (со ссылками на источник), а также:
схема оргструктуры (PNG), штатная структура (DOCX), реестр функций (XLSX).
Проекты документов требуют проверки и утверждения в установленном порядке.
"""
from __future__ import annotations

import html
import io
import re
import zipfile
from datetime import datetime
from typing import Any

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ..orgchart.render import render_chart

TYPE_RU = {"department": "департамент", "directorate": "управление", "division": "отдел", "service": "служба",
           "block": "блок", "center": "центр", "sector": "сектор", "group": "группа", "direction": "направление",
           "office": "дирекция", "branch": "филиал", "unit": "подразделение", "committee": "комитет", "lab": "лаборатория"}

L = {
    "ru": {
        "approved": "УТВЕРЖДЕНО", "by_order": "приказом {title} {company}", "date_line": "от «___» __________ 20__ г. № ____",
        "title": "ПОЛОЖЕНИЕ", "subtitle": "о структурном подразделении «{name}»",
        "s1": "1. Общие положения", "s2": "2. Основные задачи", "s3": "3. Функции", "s4": "4. Права",
        "s5": "5. Ответственность", "s6": "6. Взаимодействие", "s7": "7. Структура и штатная численность",
        "s8": "8. Заключительные положения",
        "p1_1": "Структурное подразделение «{name}» (далее — Подразделение) является структурным подразделением {company}.",
        "p1_2_parent": "Подразделение входит в состав и подчиняется руководителю структурного подразделения «{parent}».",
        "p1_2_root": "Подразделение подчиняется непосредственно руководству {company}.",
        "p1_3": "Руководство Подразделением осуществляет {head}, назначаемый на должность и освобождаемый от должности "
                "в порядке, установленном в {company}.",
        "p1_4": "В своей деятельности Подразделение руководствуется законодательством Республики Казахстан, Уставом "
                "{company}, решениями органов управления, внутренними нормативными документами и настоящим Положением.",
        "tasks_intro": "Основными задачами Подразделения являются:",
        "tasks_auto": "обеспечение выполнения функций, закреплённых за Подразделением настоящим Положением;",
        "tasks_auto2": "обеспечение эффективного взаимодействия со структурными подразделениями {company} по вопросам компетенции Подразделения.",
        "fn_intro": "Для решения возложенных задач Подразделение осуществляет следующие функции:",
        "fn_empty": "Функции Подразделения подлежат определению (в исходных документах не найдены).",
        "rights_intro": "Подразделение в пределах своей компетенции имеет право:",
        "rights_std": ["запрашивать и получать от структурных подразделений {company} информацию и документы, "
                       "необходимые для выполнения своих функций;",
                       "вносить руководству предложения по совершенствованию деятельности в пределах своей компетенции;",
                       "участвовать в совещаниях и рабочих группах по вопросам, входящим в компетенцию Подразделения;",
                       "привлекать работников других структурных подразделений (по согласованию с их руководителями) "
                       "для решения вопросов, входящих в компетенцию Подразделения."],
        "resp": ["Руководитель Подразделения несёт ответственность за ненадлежащее выполнение задач и функций, "
                 "возложенных на Подразделение настоящим Положением.",
                 "Работники Подразделения несут ответственность в соответствии с должностными инструкциями и "
                 "законодательством Республики Казахстан."],
        "inter_parent": "с руководителем структурного подразделения «{parent}» — по вопросам планирования, отчётности и организации работы;",
        "inter_child": "с подразделением «{child}» — по вопросам организации и контроля выполнения функций;",
        "inter_other": "с иными структурными подразделениями {company} — по вопросам, входящим в компетенцию Подразделения.",
        "inter_intro": "Подразделение взаимодействует:",
        "struct_sub": "В состав Подразделения входят: {items}.",
        "struct_pos": "В штат Подразделения входят должности: {items}.",
        "struct_std": "Штатная численность Подразделения определяется штатным расписанием, утверждаемым в установленном порядке.",
        "final": ["Настоящее Положение вступает в силу с даты его утверждения.",
                  "Изменения и дополнения в настоящее Положение вносятся в порядке, установленном для его утверждения."],
        "source": "источник",
        "draft": "ПРОЕКТ",
        "structure_title": "Организационная структура", "fn_registry": "Реестр функций",
        "head_default": "руководитель Подразделения",
    },
    "kz": {
        "approved": "БЕКІТІЛДІ", "by_order": "{company} {title} бұйрығымен", "date_line": "20__ жылғы «___» __________ № ____",
        "title": "«{name}»", "subtitle": "құрылымдық бөлімшесі туралы ЕРЕЖЕ",
        "s1": "1. Жалпы ережелер", "s2": "2. Негізгі міндеттер", "s3": "3. Функциялар", "s4": "4. Құқықтар",
        "s5": "5. Жауапкершілік", "s6": "6. Өзара іс-қимыл", "s7": "7. Құрылымы және штат саны",
        "s8": "8. Қорытынды ережелер",
        "p1_1": "«{name}» құрылымдық бөлімшесі (бұдан әрі – Бөлімше) {company} құрылымдық бөлімшесі болып табылады.",
        "p1_2_parent": "Бөлімше «{parent}» құрылымдық бөлімшесінің құрамына кіреді және оның басшысына бағынады.",
        "p1_2_root": "Бөлімше {company} басшылығына тікелей бағынады.",
        "p1_3": "Бөлімшеге басшылықты {company} белгіленген тәртіппен лауазымға тағайындалатын және лауазымнан "
                "босатылатын басшы ({head}) жүзеге асырады.",
        "p1_4": "Бөлімше өз қызметінде Қазақстан Республикасының заңнамасын, {company} Жарғысын, басқару органдарының "
                "шешімдерін, ішкі нормативтік құжаттарды және осы Ережені басшылыққа алады.",
        "tasks_intro": "Бөлімшенің негізгі міндеттері:",
        "tasks_auto": "осы Ережемен Бөлімшеге бекітілген функциялардың орындалуын қамтамасыз ету;",
        "tasks_auto2": "Бөлімшенің құзыреті мәселелері бойынша {company} құрылымдық бөлімшелерімен тиімді өзара іс-қимылды қамтамасыз ету.",
        "fn_intro": "Жүктелген міндеттерді шешу үшін Бөлімше мынадай функцияларды жүзеге асырады:",
        "fn_empty": "Бөлімшенің функциялары айқындалуы тиіс (бастапқы құжаттарда табылмады).",
        "rights_intro": "Бөлімше өз құзыреті шегінде:",
        "rights_std": ["{company} құрылымдық бөлімшелерінен өз функцияларын орындау үшін қажетті ақпарат пен құжаттарды "
                       "сұратуға және алуға;",
                       "басшылыққа өз құзыреті шегінде қызметті жетілдіру бойынша ұсыныстар енгізуге;",
                       "Бөлімшенің құзыретіне кіретін мәселелер бойынша кеңестер мен жұмыс топтарына қатысуға;",
                       "Бөлімшенің құзыретіне кіретін мәселелерді шешу үшін басқа құрылымдық бөлімшелердің "
                       "қызметкерлерін (олардың басшыларымен келісім бойынша) тартуға құқылы."],
        "resp": ["Бөлімше басшысы осы Ережемен Бөлімшеге жүктелген міндеттер мен функциялардың тиісінше "
                 "орындалмағаны үшін жауапты болады.",
                 "Бөлімше қызметкерлері лауазымдық нұсқаулықтарға және Қазақстан Республикасының заңнамасына сәйкес "
                 "жауапты болады."],
        "inter_parent": "«{parent}» құрылымдық бөлімшесінің басшысымен – жоспарлау, есептілік және жұмысты ұйымдастыру мәселелері бойынша;",
        "inter_child": "«{child}» бөлімшесімен – функциялардың орындалуын ұйымдастыру және бақылау мәселелері бойынша;",
        "inter_other": "{company} өзге құрылымдық бөлімшелерімен – Бөлімшенің құзыретіне кіретін мәселелер бойынша.",
        "inter_intro": "Бөлімше мыналармен өзара іс-қимыл жасайды:",
        "struct_sub": "Бөлімшенің құрамына мыналар кіреді: {items}.",
        "struct_pos": "Бөлімшенің штатына мынадай лауазымдар кіреді: {items}.",
        "struct_std": "Бөлімшенің штат саны белгіленген тәртіппен бекітілетін штат кестесімен айқындалады.",
        "final": ["Осы Ереже бекітілген күнінен бастап күшіне енеді.",
                  "Осы Ережеге өзгерістер мен толықтырулар оны бекіту үшін белгіленген тәртіппен енгізіледі."],
        "source": "дереккөз",
        "draft": "ЖОБА",
        "structure_title": "Ұйымдық құрылым", "fn_registry": "Функциялар тізілімі",
        "head_default": "Бөлімше басшысы",
    },
}


def _label(u: dict[str, Any]) -> str:
    return f"{u['name']} ({u['short']})" if u.get("short") else u["name"]


def _clean_fn(text: str) -> str:
    t = re.sub(r"^\s*[-–—•]\s*", "", text.strip())
    t = t.rstrip(" ;.,:")
    return t[:1].lower() + t[1:] if t[:2] != t[:2].upper() else t


def regulation_blocks(structure: dict[str, Any], unit_id: str, lang: str, org: dict[str, Any],
                      include_sources: bool = True, translated: dict[str, str] | None = None) -> list[dict[str, Any]]:
    """Документ как список блоков: title/center/heading/para/list."""
    t = L.get(lang, L["ru"])
    units = {u["id"]: u for u in structure["units"]}
    u = units[unit_id]
    company = org.get("company_name_kz" if lang == "kz" else "company_name") or "Общества"
    parent = units.get(u.get("parent_id") or "")
    children = [x for x in structure["units"] if x.get("parent_id") == unit_id]
    translated = translated or {}
    tr = lambda s: translated.get(s, s)  # noqa: E731

    fns = [f for f in u.get("functions", []) if f.get("kind") in ("function", "duty", "task")]
    rights = [f for f in u.get("functions", []) if f.get("kind") == "right"]
    tasks = [f for f in u.get("functions", []) if f.get("kind") == "task"]
    seen: set[str] = set()

    def items(fs: list[dict[str, Any]]) -> list[str]:
        out = []
        for f in fs:
            key = re.sub(r"\W+", "", f["text"].lower())[:120]
            if key in seen:
                continue
            seen.add(key)
            text = _clean_fn(tr(f["text"]))
            if include_sources and f.get("ref_display"):
                text += f" [{t['source']}: {f.get('doc_title') or ''} {f['ref_display']}]".replace("  ", " ")
            out.append(text + ";")
        if out:
            out[-1] = out[-1].rstrip(";") + "."
        return out

    head = u.get("head_title") or t["head_default"]
    blocks: list[dict[str, Any]] = [
        {"type": "approval", "lines": [t["approved"],
                                        t["by_order"].format(title=org.get("approver_title_kz" if lang == "kz" else "approver_title", ""),
                                                             company=company).strip(),
                                        t["date_line"]]},
        {"type": "title", "text": t["title"].format(name=_label(u))},
        {"type": "center", "text": t["subtitle"].format(name=_label(u))},
        {"type": "heading", "text": t["s1"]},
        {"type": "para", "text": "1.1. " + t["p1_1"].format(name=_label(u), company=company)},
        {"type": "para", "text": "1.2. " + (t["p1_2_parent"].format(parent=_label(parent)) if parent
                                            else t["p1_2_root"].format(company=company))},
        {"type": "para", "text": "1.3. " + t["p1_3"].format(head=head.lower() if lang == "ru" else head, company=company)},
        {"type": "para", "text": "1.4. " + t["p1_4"].format(company=company)},
        {"type": "heading", "text": t["s2"]},
        {"type": "para", "text": "2.1. " + t["tasks_intro"]},
        {"type": "list", "items": items(tasks) or [t["tasks_auto"], t["tasks_auto2"].format(company=company)]},
        {"type": "heading", "text": t["s3"]},
    ]
    fn_items = items([f for f in fns if f.get("kind") != "task"])
    if fn_items:
        blocks += [{"type": "para", "text": "3.1. " + t["fn_intro"]}, {"type": "list", "items": fn_items, "numbered": True,
                                                                       "prefix": "3.1."}]
    else:
        blocks.append({"type": "para", "text": "3.1. " + t["fn_empty"]})
    blocks += [{"type": "heading", "text": t["s4"]}, {"type": "para", "text": "4.1. " + t["rights_intro"]}]
    blocks.append({"type": "list", "items": items(rights) or [x.format(company=company) for x in t["rights_std"]]})
    blocks.append({"type": "heading", "text": t["s5"]})
    blocks += [{"type": "para", "text": f"5.{i}. {x}"} for i, x in enumerate(t["resp"], start=1)]
    inter = []
    if parent:
        inter.append(t["inter_parent"].format(parent=_label(parent)))
    inter += [t["inter_child"].format(child=_label(c)) for c in children]
    inter.append(t["inter_other"].format(company=company))
    blocks += [{"type": "heading", "text": t["s6"]}, {"type": "para", "text": "6.1. " + t["inter_intro"]},
               {"type": "list", "items": inter}]
    blocks.append({"type": "heading", "text": t["s7"]})
    n = 1
    if children:
        blocks.append({"type": "para", "text": f"7.{n}. " + t["struct_sub"].format(items=", ".join(f"«{_label(c)}»" for c in children))})
        n += 1
    positions = sorted({p["title"] for p in u.get("positions", []) if p.get("title")})
    if positions:
        blocks.append({"type": "para", "text": f"7.{n}. " + t["struct_pos"].format(items=", ".join(positions))})
        n += 1
    blocks.append({"type": "para", "text": f"7.{n}. " + t["struct_std"]})
    blocks.append({"type": "heading", "text": t["s8"]})
    blocks += [{"type": "para", "text": f"8.{i}. {x}"} for i, x in enumerate(t["final"], start=1)]
    return blocks


def blocks_to_docx(blocks: list[dict[str, Any]], draft_label: str = "") -> bytes:
    doc = DocxDocument()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for sec in doc.sections:
        sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    if draft_label:
        p = doc.add_paragraph(draft_label)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.runs[0].italic = True
    for b in blocks:
        if b["type"] == "approval":
            for ln in b["lines"]:
                p = doc.add_paragraph(ln)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.paragraph_format.space_after = Pt(0)
            doc.add_paragraph()
        elif b["type"] == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b["text"])
            r.bold = True
            r.font.size = Pt(14)
        elif b["type"] == "center":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(b["text"]).bold = True
        elif b["type"] == "heading":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.add_run(b["text"]).bold = True
        elif b["type"] == "para":
            p = doc.add_paragraph(b["text"])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)
        elif b["type"] == "list":
            for i, it in enumerate(b["items"], start=1):
                text = f"{b['prefix']}{i}. {it}" if b.get("numbered") and b.get("prefix") else f"– {it}"
                p = doc.add_paragraph(text)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = Cm(1.25)
        elif b["type"] == "image":
            doc.add_picture(io.BytesIO(b["data"]), width=Cm(16))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def blocks_to_html(blocks: list[dict[str, Any]]) -> str:
    out = ['<div class="docprev">']
    e = html.escape
    for b in blocks:
        if b["type"] == "approval":
            out.append('<div class="approval">' + "<br/>".join(e(x) for x in b["lines"]) + "</div>")
        elif b["type"] == "title":
            out.append(f"<h2>{e(b['text'])}</h2>")
        elif b["type"] == "center":
            out.append(f'<p class="center"><b>{e(b["text"])}</b></p>')
        elif b["type"] == "heading":
            out.append(f"<h3>{e(b['text'])}</h3>")
        elif b["type"] == "para":
            out.append(f"<p>{e(b['text'])}</p>")
        elif b["type"] == "list":
            tag = "ol" if b.get("numbered") else "ul"
            items = "".join(
                "<li>" + re.sub(r"\[([^\]]+)\]", r'<span class="src">[\1]</span>', e(it)) + "</li>" for it in b["items"])
            out.append(f"<{tag}>{items}</{tag}>")
    out.append("</div>")
    return "".join(out)


def _doc_titles(structure: dict[str, Any], titles: dict[int, str]) -> None:
    for u in structure["units"]:
        for f in u.get("functions", []):
            if "doc_title" not in f:
                f["doc_title"] = titles.get(f.get("doc_id"), "")


def structure_docx(structure: dict[str, Any], lang: str, org: dict[str, Any], chart_png: bytes) -> bytes:
    t = L.get(lang, L["ru"])
    company = org.get("company_name_kz" if lang == "kz" else "company_name") or ""
    units = {u["id"]: u for u in structure["units"]}
    blocks: list[dict[str, Any]] = [
        {"type": "approval", "lines": [t["approved"], t["by_order"].format(
            title=org.get("approver_title_kz" if lang == "kz" else "approver_title", ""), company=company), t["date_line"]]},
        {"type": "title", "text": f"{t['structure_title']} {company}".strip()},
        {"type": "image", "data": chart_png},
    ]
    rows = []

    def walk(uid: str | None, depth: int) -> None:
        for u in structure["units"]:
            if (u.get("parent_id") or None) == uid or (uid is None and u.get("parent_id") not in units and u.get("parent_id")):
                if u["id"] in {r[0] for r in rows}:
                    continue
                pos = ", ".join(sorted({p["title"] for p in u.get("positions", [])}))
                rows.append((u["id"], "    " * depth + _label(u) + (f" — {pos}" if pos else "")))
                walk(u["id"], depth + 1)

    walk(None, 0)
    blocks.append({"type": "list", "items": [r[1] for r in rows]})
    return blocks_to_docx(blocks)


def functions_xlsx(structure: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Функции"
    ws.append(["Подразделение", "Сокращение", "Вышестоящее", "Вид", "Функция", "Документ", "Пункт", "Общая норма"])
    units = {u["id"]: u for u in structure["units"]}
    for u in structure["units"]:
        parent = units.get(u.get("parent_id") or "")
        for f in u.get("functions", []):
            ws.append([u["name"], u.get("short", ""), _label(parent) if parent else "", f.get("kind", ""), f["text"],
                       f.get("doc_title", ""), f.get("ref_display", ""), "да" if f.get("shared") else ""])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (40, 12, 35, 14, 80, 35, 16, 12)):
        ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2 = wb.create_sheet("Штатная структура")
    ws2.append(["Подразделение", "Вышестоящее", "Руководитель", "Должности"])
    for u in structure["units"]:
        parent = units.get(u.get("parent_id") or "")
        ws2.append([_label(u), _label(parent) if parent else "", u.get("head_title", ""),
                    "; ".join(sorted({p["title"] for p in u.get("positions", [])}))])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCD", (45, 45, 25, 70)):
        ws2.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def slug(text: str) -> str:
    s = re.sub(r"[^\w\s-]", "", text, flags=re.UNICODE).strip()
    return re.sub(r"\s+", "_", s)[:60] or "unit"


def build_package(structure: dict[str, Any], org: dict[str, Any], langs: list[str], unit_ids: list[str] | None,
                  doc_titles: dict[int, str], include_sources: bool = True,
                  translations: dict[str, str] | None = None, extra_files: dict[str, bytes] | None = None) -> bytes:
    _doc_titles(structure, doc_titles)
    chart_units = [{"id": u["id"], "name": u["name"], "short": u.get("short"), "parent_id": u.get("parent_id")}
                   for u in structure["units"]]
    chart = render_chart(chart_units, title=org.get("company_name", ""))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("00_Схема_оргструктуры.png", chart)
        zf.writestr("01_Реестр_функций_и_штатная_структура.xlsx", functions_xlsx(structure))
        for lang in langs:
            suffix = "KZ" if lang == "kz" else "RU"
            zf.writestr(f"02_Организационная_структура_{suffix}.docx", structure_docx(structure, lang, org, chart))
            targets = unit_ids or [u["id"] for u in structure["units"]]
            for i, uid in enumerate(targets, start=1):
                u = next((x for x in structure["units"] if x["id"] == uid), None)
                if not u:
                    continue
                blocks = regulation_blocks(structure, uid, lang, org, include_sources,
                                           translations if lang == "kz" else None)
                name = f"Положения_{suffix}/{i:02d}_Положение_{slug(u.get('short') or u['name'])}_{suffix}.docx"
                zf.writestr(name, blocks_to_docx(blocks, L[lang]["draft"]))
        for name, data in (extra_files or {}).items():
            zf.writestr(name, data)
        zf.writestr("README.txt", (
            "Пакет документов сформирован автоматически " + datetime.now().strftime("%d.%m.%Y %H:%M") + ".\n"
            "Документы являются ПРОЕКТАМИ и требуют проверки ответственным сотрудником и утверждения\n"
            "в установленном порядке. В квадратных скобках указаны источники функций (документ и пункт).\n"
        ).encode("utf-8"))
    return buf.getvalue()
