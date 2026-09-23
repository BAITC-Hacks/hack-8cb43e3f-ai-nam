"""Экспорт результатов: заключение (DOCX/PDF) и таблица сопоставления функций (XLSX)."""
from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..analysis.compare import FN_STATUS_LABELS
from ..analysis.conclusion import T, build_conclusion, ref_for
from ..analysis.recommend import recommendation_text
from .xlsx import literal_cells

SEV_RU = {"high": "высокая", "medium": "средняя", "low": "низкая", "info": "инфо"}
SEV_KZ = {"high": "жоғары", "medium": "орташа", "low": "төмен", "info": "ақпарат"}
SEV_COLOR = {"high": "C0392B", "medium": "D68910", "low": "2E86C1", "info": "7F8C8D"}
REVIEW_RU = {"confirmed": "подтверждено", "rejected": "отклонено", "pending": "ожидает проверки"}
REVIEW_KZ = {"confirmed": "расталды", "rejected": "қабылданбады", "pending": "тексеруді күтуде"}
FN_STATUS_KZ = {"preserved": "сақталған", "transferred": "берілген", "modified": "өзгертілген", "lost": "жоғалған",
                "relocated": "жалпы ережелерге көшірілген", "new": "жаңа"}


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _base_doc() -> DocxDocument:
    doc = DocxDocument()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for sec in doc.sections:
        sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    return doc


def conclusion_docx(conclusion: dict[str, Any], result: dict[str, Any], lang: str = "ru",
                    org: dict[str, Any] | None = None) -> bytes:
    t = T.get(lang, T["ru"])
    sev = SEV_KZ if lang == "kz" else SEV_RU
    rev = REVIEW_KZ if lang == "kz" else REVIEW_RU
    doc = _base_doc()
    org = org or {}
    company = org.get("company_name_kz" if lang == "kz" else "company_name") or ""
    if company:
        p = doc.add_paragraph(company)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(conclusion["title"])
    r.bold = True
    r.font.size = Pt(14)
    date = conclusion.get("generated_at", "")[:10]
    p = doc.add_paragraph(("Дата формирования: " if lang == "ru" else "Қалыптастырылған күні: ") + date)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(10)

    for sec in conclusion["sections"]:
        hp = doc.add_paragraph()
        hr = hp.add_run(sec["title"])
        hr.bold = True
        hr.font.size = Pt(13)
        hp.paragraph_format.space_before = Pt(12)
        for para in sec.get("paragraphs", []):
            if para:
                pp = doc.add_paragraph(para)
                pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if sec.get("method"):
            mp = doc.add_paragraph(sec["method"])
            mp.runs[0].italic = True
            mp.runs[0].font.size = Pt(9)
        for i, it in enumerate(sec.get("items", []), start=1):
            ip = doc.add_paragraph(style="List Number" if sec["id"] not in ("scope",) else None)
            if it.get("finding_id"):
                tag = ip.add_run(f"[{it['finding_id']}] ")
                tag.bold = True
                tag.font.color.rgb = RGBColor.from_string(SEV_COLOR.get(it.get("severity", "info"), "7F8C8D"))
            ip.add_run(it["text"])
            meta = []
            if it.get("severity") and it.get("finding_id"):
                meta.append(("Значимость: " if lang == "ru" else "Маңыздылығы: ") + sev.get(it["severity"], ""))
            if it.get("finding_id"):
                meta.append(("Проверка: " if lang == "ru" else "Тексеру: ") + rev.get(it.get("review", "pending"), ""))
            if it.get("disputed"):
                meta.append("оспорено ИИ-моделью" if lang == "ru" else "ЖИ-модель дау айтты")
            if meta:
                mr = ip.add_run("  (" + "; ".join(meta) + ")")
                mr.font.size = Pt(9)
                mr.italic = True
            for src in it.get("sources", [])[:4]:
                sp = doc.add_paragraph()
                sp.paragraph_format.left_indent = Cm(1)
                sr = sp.add_run(f"{t['sources']}: {src}")
                sr.italic = True
                sr.font.size = Pt(9)
                sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Приложение: статусы подразделений
    doc.add_page_break()
    ap = doc.add_paragraph()
    ar = ap.add_run("Приложение 1. Статусы подразделений" if lang == "ru" else "1-қосымша. Бөлімшелердің мәртебелері")
    ar.bold = True
    head = ["До", "После", "Статус", "Функций (до)", "Не найдено"] if lang == "ru" else \
        ["Дейін", "Кейін", "Мәртебе", "Функциялар (дейін)", "Табылмады"]
    table = doc.add_table(rows=1, cols=len(head))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hname in enumerate(head):
        table.rows[0].cells[i].text = hname
        _shade(table.rows[0].cells[i], "D9E2F3")
    for u in result["units"]:
        row = table.add_row().cells
        row[0].text = (u["before"]["name"] + (f" ({u['before']['short']})" if u["before"].get("short") else "")) if u["before"] else "—"
        row[1].text = ", ".join(a["name"] + (f" ({a['short']})" if a.get("short") else "") for a in u["after"]) or "—"
        row[2].text = t["status"].get(u["status"], u["status"])
        row[3].text = str(u.get("functions_total") or "")
        row[4].text = str(u.get("functions_lost") or "")
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for rr in p.runs:
                    rr.font.size = Pt(9)

    # Приложение: изменённые функции
    doc.add_paragraph()
    ap = doc.add_paragraph()
    ar = ap.add_run("Приложение 2. Сопоставление функций (изменения)" if lang == "ru"
                    else "2-қосымша. Функцияларды салыстыру (өзгерістер)")
    ar.bold = True
    head = ["Статус", "Функция «до» (источник)", "Соответствие «после» (источник)", "Сходство"] if lang == "ru" else \
        ["Мәртебе", "«Дейінгі» функция (дереккөз)", "«Кейінгі» сәйкестік (дереккөз)", "Ұқсастық"]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, hname in enumerate(head):
        table.rows[0].cells[i].text = hname
        _shade(table.rows[0].cells[i], "D9E2F3")
    labels = FN_STATUS_KZ if lang == "kz" else FN_STATUS_LABELS
    for row in result["function_map"]:
        if row["status"] == "preserved":
            continue
        cells = table.add_row().cells
        cells[0].text = labels.get(row["status"], row["status"])
        b = row.get("before")
        a = row["after"][0] if row.get("after") else None
        cells[1].text = f"{b['unit']}: {b['text'][:300]}\n({ref_for(b, lang)})" if b else "—"
        cells[2].text = f"{a['unit']}: {a['text'][:300]}\n({ref_for(a, lang)})" if a and row["status"] != "lost" else "—"
        cells[3].text = f"{row['score']:.2f}"
        if row["status"] == "lost":
            _shade(cells[0], "F5B7B1")
        elif row["status"] == "new":
            _shade(cells[0], "ABEBC6")
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for rr in p.runs:
                    rr.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


XLSX = {
    "ru": {
        "sheet_map": "Сопоставление функций", "sheet_units": "Подразделения", "sheet_findings": "Выводы",
        "map_head": ["ID", "Статус", "Сходство", "Подразделение «до»", "Функция «до»", "Источник «до»",
                     "Подразделение «после»", "Функция «после»", "Источник «после»", "Проверка ИИ", "Вывод"],
        "units_head": ["Статус", "До", "После", "Функций (до)", "Не найдено", "Покрытие в том же подразделении",
                       "Куда переданы"],
        "findings_head": ["ID", "Тип", "Значимость", "Вывод", "Подразделения", "Источники", "Рекомендация", "Метод",
                          "Комментарий ИИ"],
        "types": {"loss": "Потеря функции", "duplication": "Дублирование", "conflict": "Конфликт интересов",
                  "reorganization": "Изменение структуры", "improvement": "Улучшение",
                  "requirement_gap": "Пробел в требованиях", "benchmark": "Практика операторов"},
    },
    "kz": {
        "sheet_map": "Функцияларды салыстыру", "sheet_units": "Бөлімшелер", "sheet_findings": "Қорытындылар",
        "map_head": ["ID", "Мәртебе", "Ұқсастық", "«Дейінгі» бөлімше", "«Дейінгі» функция", "«Дейінгі» дереккөз",
                     "«Кейінгі» бөлімше", "«Кейінгі» функция", "«Кейінгі» дереккөз", "ЖИ тексеруі", "Қорытынды"],
        "units_head": ["Мәртебе", "Дейін", "Кейін", "Функциялар (дейін)", "Табылмады", "Сол бөлімшедегі қамту",
                       "Қайда берілді"],
        "findings_head": ["ID", "Түрі", "Маңыздылығы", "Қорытынды", "Бөлімшелер", "Дереккөздер", "Ұсыным", "Әдіс",
                          "ЖИ түсіндірмесі"],
        "types": {"loss": "Функцияның жоғалуы", "duplication": "Қайталану", "conflict": "Мүдделер қақтығысы",
                  "reorganization": "Құрылымның өзгеруі", "improvement": "Жақсару",
                  "requirement_gap": "Талаптардағы олқылық", "benchmark": "Операторлар тәжірибесі"},
    },
}


def function_map_xlsx(result: dict[str, Any], lang: str = "ru") -> bytes:
    lang = lang if lang in XLSX else "ru"
    x = XLSX[lang]
    fn_labels = FN_STATUS_KZ if lang == "kz" else FN_STATUS_LABELS
    sev = SEV_KZ if lang == "kz" else SEV_RU
    # формулировки выводов на казахском берём из заключения (шаблоны T["kz"]); цитаты документов не переводятся
    titles: dict[str, str] = {}
    if lang != "ru":
        concl = build_conclusion(result, result.get("documents", []), lang)
        for sec in concl["sections"]:
            if sec["id"] == "recommendations":
                continue
            for it in sec["items"]:
                if it.get("finding_id"):
                    titles.setdefault(it["finding_id"], it["text"])

    wb = Workbook()
    ws = wb.active
    ws.title = x["sheet_map"]
    ws.append(x["map_head"])
    fills = {"lost": "F5B7B1", "new": "ABEBC6", "transferred": "FAD7A0", "modified": "F9E79F", "relocated": "D6EAF8"}
    for row in result["function_map"]:
        b = row.get("before") or {}
        a = row["after"][0] if row.get("after") else {}
        verified = row.get("verified") or {}
        ws.append([
            row["id"], fn_labels.get(row["status"], row["status"]), row["score"],
            b.get("unit", ""), b.get("text", ""), ref_for(b, lang) if b else "",
            a.get("unit", "") if row["status"] != "lost" else "", a.get("text", "") if row["status"] != "lost" else "",
            ref_for(a, lang) if a and row["status"] != "lost" else "",
            f"{verified.get('verdict', '')}: {verified.get('reason', '')}" if verified else "",
            row.get("finding_id", ""),
        ])
        if row["status"] in fills:
            ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=fills[row["status"]])
    widths = [8, 16, 10, 30, 60, 30, 30, 60, 30, 30, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E2F3")
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet(x["sheet_units"])
    ws2.append(x["units_head"])
    for u in result["units"]:
        ws2.append([
            T[lang]["status"].get(u["status"], u["status"]),
            (u["before"]["name"] + (f" ({u['before']['short']})" if u["before"].get("short") else "")) if u["before"] else "",
            ", ".join(a["name"] for a in u["after"]),
            u.get("functions_total"), u.get("functions_lost"), u.get("coverage"),
            "; ".join(f"{d['unit']} ({d['count']})" for d in u.get("destinations", [])),
        ])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for i, w in enumerate([18, 45, 45, 12, 12, 16, 70], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet(x["sheet_findings"])
    ws3.append(x["findings_head"])
    for f in result["findings"]:
        rec = (recommendation_text(f, result["units"], lang) if lang != "ru" else "") or f.get("recommendation", "")
        ws3.append([
            f["id"], x["types"].get(f["type"], f["type"]), sev.get(f["severity"], f["severity"]),
            titles.get(f["id"], f["title"]), ", ".join(f.get("units", [])),
            "\n".join(ref_for(e, lang) for e in f.get("evidence", [])[:5]), rec,
            f.get("method", ""), f.get("llm_note", ""),
        ])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for i, w in enumerate([8, 16, 12, 70, 40, 50, 60, 12, 40], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    for r in ws3.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    literal_cells(wb).save(buf)
    return buf.getvalue()


def soffice_available() -> bool:
    return bool(shutil.which("soffice") or shutil.which("libreoffice"))


def docx_to_pdf(docx_bytes: bytes) -> bytes | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "doc.docx"
        src.write_bytes(docx_bytes)
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                           check=True, timeout=180, capture_output=True,
                           env={"HOME": tmp, "PATH": "/usr/bin:/bin:/usr/local/bin"})
        except (subprocess.SubprocessError, OSError):
            return None
        out = Path(tmp) / "doc.pdf"
        return out.read_bytes() if out.exists() else None
