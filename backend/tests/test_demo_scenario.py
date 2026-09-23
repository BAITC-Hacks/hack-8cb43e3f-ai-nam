"""Контрольный сценарий из ТЗ (п. 11): реорганизация, потеря функции, дублирование + конфликт интересов.

Комплект samples/demo генерируется scripts/make_demo.py; ожидаемые изменения описаны в samples/README.md.
"""
from pathlib import Path

import pytest

from app.analysis.compare import Comparator
from app.analysis.extras import check_requirements
from app.analysis.structure import build_structure
from app.analysis.text import Similarity
from app.parsing.parse import parse_file
from app.services.settings_store import ANALYSIS_DEFAULTS

DEMO = Path(__file__).resolve().parents[2] / "samples" / "demo"


@pytest.fixture(scope="module")
def result():
    docs, i = {}, 0
    for side in ("before", "after", "requirements"):
        docs[side] = []
        for f in sorted((DEMO / side).iterdir()):
            i += 1
            docs[side].append({"id": i, "filename": f.name, "title": f.stem, "parsed": parse_file(f)})
    kinds = ANALYSIS_DEFAULTS["function_kinds"]
    before = build_structure("before", docs["before"], kinds)
    after = build_structure("after", docs["after"], kinds)
    cmp = Comparator(before, after, ANALYSIS_DEFAULTS, Similarity(), {d["id"]: d["title"] for s in docs.values() for d in s})
    res = cmp.run()
    res["requirements"] = check_requirements(cmp, docs["requirements"])
    return res


def _status(res, name):
    for u in res["units"]:
        names = [(u["before"] or {}).get("name", "")] + [a["name"] for a in u["after"]]
        if any(name in n for n in names):
            yield u["status"]


def test_reorganization_detected(result):
    assert "split" in set(_status(result, "Департамент информационных технологий"))
    assert "created" in set(_status(result, "Департамент цифровой инфраструктуры"))
    assert "created" in set(_status(result, "Служба информационной безопасности"))
    assert "renamed" in set(_status(result, "Департамент внутреннего контроля и рисков"))
    assert "preserved" in set(_status(result, "Департамент закупок"))


def test_lost_function_with_source(result):
    losses = [f for f in result["findings"] if f["type"] == "loss"]
    it = [f for f in losses if "ИТ-стратегии" in f["title"]]
    assert it, "потеря функции «ИТ-стратегия» не обнаружена"
    ev = it[0]["evidence"][0]
    assert ev["side"] == "before" and ev["ref_display"] == "п. 3.1" and "ДИТ" in ev["doc_title"]


def test_duplication_detected(result):
    dups = [f for f in result["findings"] if f["type"] == "duplication"]
    assert any("мониторинг исполнения договоров" in f["title"] for f in dups)
    units = {u for f in dups for u in f["units"]}
    assert any("закупок" in u for u in units) and any("Юридический" in u for u in units)


def test_conflict_of_interest_detected(result):
    conf = [f for f in result["findings"] if f["type"] == "conflict" and f.get("subtype") == "sod"]
    assert any("СИБ" in " ".join(f["units"]) for f in conf)


def test_requirement_gap(result):
    gaps = [f for f in result["findings"] if f["type"] == "requirement_gap"]
    assert any("ИТ-стратегии" in f["title"] for f in gaps)


def test_every_finding_has_source(result):
    for f in result["findings"]:
        assert f["evidence"], f"вывод {f['id']} без источника"
        for e in f["evidence"]:
            assert e.get("doc_id") and e.get("ref_display") is not None


def test_kz_recommendations_and_xlsx(result):
    import copy
    from io import BytesIO

    from openpyxl import load_workbook

    from app.analysis.conclusion import build_conclusion
    from app.analysis.recommend import add_recommendations
    from app.docgen.report import FN_STATUS_KZ, function_map_xlsx

    res = copy.deepcopy(result)
    add_recommendations(res)
    ru = {f["id"]: f["recommendation"] for f in res["findings"] if f.get("recommendation")}
    recs = next(s for s in build_conclusion(res, [], "kz")["sections"] if s["id"] == "recommendations")
    assert recs["items"]
    for it in recs["items"]:
        assert it["text"] != ru[it["finding_id"]]
        assert any(ch in it["text"] for ch in "әіңғүұқөһ"), it["text"]

    wb = load_workbook(BytesIO(function_map_xlsx(res, "kz")))
    assert wb.sheetnames == ["Функцияларды салыстыру", "Бөлімшелер", "Қорытындылар"]
    statuses = {r[1] for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True)}
    assert statuses <= set(FN_STATUS_KZ.values())


def test_xlsx_writes_document_text_literally(result):
    import copy
    from io import BytesIO

    from openpyxl import load_workbook

    from app.docgen.report import function_map_xlsx

    res = copy.deepcopy(result)
    formula = '=HYPERLINK("http://example.com","x")'
    res["function_map"][0]["before"]["text"] = formula
    cell = load_workbook(BytesIO(function_map_xlsx(res))).worksheets[0]["E2"]
    assert cell.value == formula and cell.data_type == "s"

