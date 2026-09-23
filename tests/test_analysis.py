import io
from pathlib import Path
import unittest

from app.parsers import parse_document
from app.analysis import analyze
from app.reports import html_report, xlsx_report

ROOT = Path(__file__).resolve().parents[1]


def textdoc(text, side):
    return parse_document(text.encode('utf-8'), side + '.txt', side, side)


class AnalysisTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.demo_docs = [parse_document((ROOT / 'samples' / f'audit-v{v}.docx').read_bytes(), f'audit-v{v}.docx', side, side) for v, side in [('8', 'before'), ('9', 'after')]]
        cls.demo = analyze(cls.demo_docs)
        cls.refs = {c.id: c for d in cls.demo_docs for c in d.clauses}

    def test_real_department_changes(self):
        self.assertEqual(len(self.demo['structure']), 4)
        added = [s['name'] for s in self.demo['structure'] if s['status'] == 'added']
        self.assertEqual(len(added), 2)
        self.assertTrue(any('ДИТААД' in s for s in added))
        self.assertTrue(any('ДОА' in s for s in added))

    def test_transfer_is_not_loss_even_for_parent_list(self):
        rows = [m for m in self.demo['matrix'] if m['before_ids'] and self.refs[m['before_ids'][0]].number.startswith('5.4.4')]
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(m['kind'] == 'transferred' for m in rows), rows)
        self.assertEqual({self.refs[m['after_ids'][0]].number for m in rows}, {'5.3.3', '5.3.3.а', '5.3.3.б'})

    def test_obligation_change_at_9_15(self):
        rows = [m for m in self.demo['matrix'] if m['before_ids'] and self.refs[m['before_ids'][0]].number == '9.15']
        self.assertEqual(rows[0]['kind'], 'modality')

    def test_specific_power_missing(self):
        rows = [m for m in self.demo['matrix'] if m['before_ids'] and self.refs[m['before_ids'][0]].number == '5.6.2']
        self.assertEqual(rows[0]['kind'], 'missing')

    def test_all_findings_have_real_sources(self):
        for f in self.demo['findings']:
            ids = f['before_ids'] + f['after_ids']
            self.assertTrue(ids)
            self.assertTrue(all(cid in self.refs for cid in ids))
            self.assertTrue(all(self.refs[cid].doc_id == 'before' for cid in f['before_ids']))
            self.assertTrue(all(self.refs[cid].doc_id == 'after' for cid in f['after_ids']))

    def test_identical_documents_do_not_create_risks(self):
        text = '1. Функции\n1.1. Директор отдела:\n1.1.1. Организует проверку информационных систем.\n1.1.2. Готовит отчет о проверке.'
        result = analyze([textdoc(text,'before'), textdoc(text,'after')])
        self.assertFalse(result['findings'])
        self.assertEqual(result['stats']['coverage'], 100)

    def test_number_change_preserves_function(self):
        old = '1. Функции\n1.1. Организует контроль качества закупок.'
        new = '1. Функции\n1.2. Организует контроль качества закупок.'
        result = analyze([textdoc(old,'before'), textdoc(new,'after')])
        self.assertEqual(result['matrix'][0]['kind'], 'renumbered')
        self.assertFalse(any(f['kind'] == 'missing' for f in result['findings']))

    def test_lost_function_is_reported_with_qualification(self):
        old = '1. Функции\n1.1. Организует проверку пожарной безопасности объектов.\n1.2. Готовит бухгалтерскую отчетность.'
        new = '1. Функции\n1.2. Готовит бухгалтерскую отчетность.'
        result = analyze([textdoc(old,'before'),textdoc(new,'after')])
        missing = [f for f in result['findings'] if f['kind']=='missing']
        self.assertEqual(len(missing), 1)
        self.assertIn('не доказательство', missing[0]['description'])

    def test_negation_is_not_editorial(self):
        old = '1. Функции\n1.1. Руководитель осуществляет проверку всех закупок компании.'
        new = '1. Функции\n1.1. Руководитель не осуществляет проверку всех закупок компании.'
        result = analyze([textdoc(old,'before'),textdoc(new,'after')])
        self.assertEqual(result['matrix'][0]['kind'],'modified')

    def test_xlsx_locations(self):
        from openpyxl import Workbook
        book=Workbook();book.active.title='Функции';book.active.append(['1.1.', 'Организует проверку закупок']);stream=io.BytesIO();book.save(stream)
        doc=parse_document(stream.getvalue(),'source.xlsx','before','sheet')
        self.assertIn('A1:B1',doc.clauses[0].location)
        self.assertIn('Функции',doc.clauses[0].location)

    def test_scanned_pdf_is_not_silently_analyzed(self):
        from pypdf import PdfWriter
        book=PdfWriter();book.add_blank_page(595,842);stream=io.BytesIO();book.write(stream)
        with self.assertRaisesRegex(ValueError,'OCR'):
            parse_document(stream.getvalue(),'scan.pdf','before','scan')

    def test_unsupported_format(self):
        with self.assertRaises(ValueError):parse_document(b'content','file.exe','before','bad')

    def test_consolidated_identical_duties_are_not_lost(self):
        before='1. Функции\n1.1. Готовит отчет по закупкам.\n1.2. Готовит отчет по закупкам.'
        after='1. Функции\n1.1. Готовит отчет по закупкам.'
        result=analyze([textdoc(before,'before'),textdoc(after,'after')])
        self.assertEqual(result['stats']['matched'],2)
        self.assertFalse(any(f['kind']=='missing' for f in result['findings']))

    def test_docx_automatic_numbering_and_tables(self):
        from docx import Document
        doc=Document();doc.add_paragraph('1. Функции');doc.add_paragraph('1.1. Организует проверку информационной безопасности.');table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='2.1. Готовит отчеты об устранении нарушений.'
        stream=io.BytesIO();doc.save(stream);parsed=parse_document(stream.getvalue(),'test.docx','before','word')
        self.assertTrue(any(c.number=='2.1' and 'Таблица' in c.location for c in parsed.clauses))

    def test_report_is_standalone_and_escaped(self):
        result=analyze([textdoc('1. Функции\n1.1. Готовит <script>alert(1)</script> отчет.', 'before'),textdoc('1. Функции\n1.1. Готовит отчет.', 'after')])
        case={'title':'<img onerror=alert(1)>','owner_name':'Reviewer','result':result}
        report=html_report(case)
        self.assertNotIn('<script>',report)
        self.assertIn('&lt;script&gt;',report)
        for f in result['findings']:
            for cid in f['before_ids']+f['after_ids']:self.assertIn('id="'+cid+'"',report)
        from openpyxl import load_workbook
        xlsx=load_workbook(io.BytesIO(xlsx_report(case)))
        self.assertEqual(xlsx.sheetnames,['Выводы','Матрица функций','Источники','Подразделения'])

    def test_word_is_editable_with_verbatim_evidence_and_review(self):
        from docx import Document
        from app.reports import docx_report
        result=analyze([textdoc('1. Функции\n1.1. Организует проверку закупок.', 'before'),textdoc('1. Функции\n1.1. Может организовать проверку закупок.', 'after')])
        finding=result['findings'][0]
        finding.update(review='confirmed',comment='Сверено с оригиналом',reviewer='Reviewer')
        doc=Document(io.BytesIO(docx_report({'title':'Контроль Word','owner_name':'Reviewer','result':result})))
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Сверено с оригиналом',text)
        self.assertIn('Подтверждено',text)
        self.assertIn('1.1. Организует проверку закупок.',text)
        self.assertIn('1.1. Может организовать проверку закупок.',text)
        self.assertTrue(doc.tables)
        self.assertEqual(doc.paragraphs[0].style.name,'Title')


if __name__=='__main__': unittest.main()
